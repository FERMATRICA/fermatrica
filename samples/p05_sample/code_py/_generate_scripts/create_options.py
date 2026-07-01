"""
Create options.xlsx and seasonality.xlsx for both BHT and Sales models.
All files go into the shared `model_data/options/` directory.

Options.xlsx structure (per-media-channel format matching p01/p02):
  - translation_1  → BHT model item→variable mappings
  - translation_2  → Sales model item→variable mappings
  - main           → budget scenarios (predefined, zero, etc.)
  - current        → current budget scenario

Seasonality.xlsx shared between both nodes with monthly sin-wave values.
"""
import os
import numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
p05_root = os.path.join(_SCRIPT_DIR, '..', '..')

OUT_DIR = os.path.join(p05_root, 'code_py', 'model_data', 'options')


def write_sheet(ws, headers, rows, col_widths=None):
    for ci, h in enumerate(headers, 1):
        ws.cell(row=1, column=ci, value=h)
    for ri, row in enumerate(rows, 2):
        for ci, val in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=val)
    if col_widths:
        for ci, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w


# ── Media channel definitions ──────────────────────────────────
#   (budget_item, variable, max_budg, predefined_value)
MEDIA_CHANNELS = [
    ('bdg_tv',          'brnd_tv_ots_clip',         500, 200),
    ('bdg_ooh',         'brnd_ots_ooh',             100,  30),
    ('bdg_dgt_perform',  'brnd_cl_dgtind_perform',   50,  20),
    ('bdg_dgt_display',  'brnd_cl_dgtind_display',  100,  30),
    ('bdg_dgt_video',    'brnd_cl_dgtind_video',    100,  30),
    ('bdg_dgt_special',  'brnd_cl_dgtind_special',  100,  20),
    ('bdg_dgt_indoor',   'brnd_cl_dgtind_indoor',    70,  15),
    ('bdg_ecom',         'e_com',                    30,  10),
    ('bdg_pr',           'pr_costs_sm',              30,  10),
    ('bdg_pos',          'pos',                      50,  20),
]

BUDGET_ITEMS = [ch[0] for ch in MEDIA_CHANNELS]
BUDGET_COLS = ['option', 'text_flags'] + BUDGET_ITEMS

# Seasonality columns = budget items stripped of bdg_/nobdg_ prefix
# (media_ssn_apply will prefix these with opt_media_ssn_)
def _ssn_col(budget_item: str) -> str:
    for pfx in ('bdg_', 'nobdg_'):
        if budget_item.startswith(pfx):
            return budget_item[len(pfx):]
    return budget_item

SSN_COLS = [_ssn_col(it) for it in BUDGET_ITEMS]


# ════════════════════════════════════════════════════════════════
#   options.xlsx
# ════════════════════════════════════════════════════════════════
out_opts = os.path.join(OUT_DIR, 'options.xlsx')
wb = Workbook()

# ── translation_1 (BHT model) ──
ws_t1 = wb.active
ws_t1.title = 'translation_1'
trans_headers = ['item', 'function', 'variable', 'max_budg', 'max_obs_budg']
trans_rows = [[item, 'basic', var, mx, round(mx / 1.35)] for item, var, mx, _ in MEDIA_CHANNELS]
write_sheet(ws_t1, trans_headers, trans_rows, [25, 12, 30, 10])

# ── translation_2 (Sales model) ──
ws_t2 = wb.create_sheet('translation_2')
write_sheet(ws_t2, trans_headers, trans_rows, [25, 12, 30, 10])

# ── main ──
ws_main = wb.create_sheet('main')
main_rows = [
    ['zero',       None] + [0] * len(BUDGET_ITEMS),
    ['predefined', None] + [ch[3] for ch in MEDIA_CHANNELS],
]
write_sheet(ws_main, BUDGET_COLS, main_rows, [15, 12] + [12] * len(BUDGET_ITEMS))

# ── current ──
ws_cur = wb.create_sheet('current')
cur_rows = [
    ['predefined', None] + [ch[3] for ch in MEDIA_CHANNELS],
]
write_sheet(ws_cur, BUDGET_COLS, cur_rows, [15, 12] + [12] * len(BUDGET_ITEMS))

wb.save(out_opts)
print(f'Created: {out_opts}')


# ════════════════════════════════════════════════════════════════
#   seasonality.xlsx  (shared, monthly sin-wave, per-channel)
#   Each media channel gets its own column (bare name without
#   bdg_/nobdg_ prefix).  media_ssn_apply auto-prefixes these
#   with 'opt_media_ssn_' during merge.
# ════════════════════════════════════════════════════════════════
out_ssn = os.path.join(OUT_DIR, 'seasonality.xlsx')
wb = Workbook()
ws = wb.active
ws.title = 'data'
ssn_headers = ['month'] + SSN_COLS
raw = [0.5 + 0.4 * np.sin(2 * np.pi * (m - 1) / 12) for m in range(1, 13)]
total = sum(raw)
ssn_rows = []
for m in range(1, 13):
    val = round(raw[m - 1] / total, 6)
    ssn_rows.append([m] + [val] * len(SSN_COLS))
write_sheet(ws, ssn_headers, ssn_rows, [10] + [20] * len(SSN_COLS))
wb.save(out_ssn)
print(f'Created: {out_ssn}')


print('\nAll options and seasonality files created.')