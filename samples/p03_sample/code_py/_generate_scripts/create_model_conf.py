import os
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


def write_sheet(ws, headers, rows, col_widths):
    for ci, h in enumerate(headers, 1):
        ws.cell(row=1, column=ci, value=h)
    for ri, row in enumerate(rows, 2):
        for ci, val in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=val)
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w


# Load data to compute initial values
dt = pd.read_excel(
    os.path.join(_SCRIPT_DIR, '..', '..', 'data', 'data_processed', 'data_for_sample_p03.xlsx'),
    parse_dates=['date'])
train_mask = dt['listed'].isin([2])


def train_mean(var):
    m = (dt[var] > 0.0) & train_mask
    if m.sum() == 0:
        return dt.loc[train_mask, var].mean()
    return dt.loc[m, var].mean()


def train_std(var):
    m = (dt[var] > 0.0) & train_mask
    if m.sum() == 0:
        return dt.loc[train_mask, var].std()
    return dt.loc[m, var].std()


# Compute stats
tv_avg = train_mean('brnd_tv_ots_clip')
tv_std = train_std('brnd_tv_ots_clip')
ooh_avg = train_mean('brnd_ots_ooh')
ooh_std = train_std('brnd_ots_ooh')
dgt_perf_avg = train_mean('brnd_cl_dgtind_perform')
dgt_perf_std = train_std('brnd_cl_dgtind_perform')
dgt_disp_avg = train_mean('brnd_cl_dgtind_display')
dgt_disp_std = train_std('brnd_cl_dgtind_display')
dgt_vid_avg = train_mean('brnd_cl_dgtind_video')
dgt_vid_std = train_std('brnd_cl_dgtind_video')
dgt_spec_avg = train_mean('brnd_cl_dgtind_special')
dgt_spec_std = train_std('brnd_cl_dgtind_special')
dgt_indoor_avg = train_mean('brnd_cl_dgtind_indoor')
dgt_indoor_std = train_std('brnd_cl_dgtind_indoor')
ecom_avg = train_mean('e_com')
ecom_std = train_std('e_com')
pos_avg = train_mean('pos')
pos_std = train_std('pos')
pr_avg = train_mean('pr_costs_sm')
pr_std = train_std('pr_costs_sm')
aba_lag_avg = train_mean('aba_lag')
aba_lag_std = train_std('aba_lag')

price_mar_n = 17
price_mar_avg = train_mean('price_ret_rel_market')
price_mar_std = train_std('price_ret_rel_market')

print(f'TV: avg={tv_avg:.2f}, std={tv_std:.2f}')
print(f'OOH: avg={ooh_avg:.2f}, std={ooh_std:.2f}')
print(f'Digital perf: avg={dgt_perf_avg:.2f}, std={dgt_perf_std:.2f}')
print(f'Digital disp: avg={dgt_disp_avg:.2f}, std={dgt_disp_std:.2f}')
print(f'Digital video: avg={dgt_vid_avg:.2f}, std={dgt_vid_std:.2f}')
print(f'Digital special: avg={dgt_spec_avg:.2f}, std={dgt_spec_std:.2f}')
print(f'Digital indoor: avg={dgt_indoor_avg:.2f}, std={dgt_indoor_std:.2f}')
print(f'E-com: avg={ecom_avg:.2f}, std={ecom_std:.2f}')
print(f'POS: avg={pos_avg:.2f}, std={pos_std:.2f}')
print(f'PR: avg={pr_avg:.2f}, std={pr_std:.2f}')
print(f'ABA lag: avg={aba_lag_avg:.2f}, std={aba_lag_std:.2f}')
print(f'Price rel mkt: avg={price_mar_avg:.2f}, std={price_mar_std:.2f}')

# Set avg upper bound as 5x the actual avg
def avg_upper(v):
    return max(v * 5, v + 1)


out_path = os.path.join(_SCRIPT_DIR, '..', 'model_data', 'model_conf.xlsx')
wb = Workbook()

# ═══════════════════════════════════════════
# Sheet: params
# ═══════════════════════════════════════════
params_headers = ['variable', 'fun', 'arg', 'lower', 'upper', 'value', 'type', 'fixed', 'if_active',
                  'index_vars', 'index_aggr', 'index_free_var', 'set_start', 'formula']

FM = 'float64'

params_rows = [
    # === Distribution: di_soph_ssn_median_sm ===
    ['di_soph_ssn_median_sm', 'softmax', 'lambda', 5.0, 40.0, 33.446043, FM, 0, 1],
    ['di_soph_ssn_median_sm', 'softmax', 'avg', 0.0, 100.0, 37.148661, FM, 0, 1],
    ['di_soph_ssn_median_sm', 'softmax', 'std', None, None, 0.128714, FM, 1, 1, None, None, None, 1,
     '''ds.loc[(ds["di_soph_ssn_median_sm"] > 0.0) & (ds["listed"].isin([2])), "di_soph_ssn_median_sm"].std()'''],
    ['di_soph_ssn_median_sm', 'share', 'coef_var', None, None, None, 'str', 1, 1, '"market","date","listed"'],
    ['di_soph_ssn_median_sm', 'share', 'include_coef_var', None, None, False, 'bool', 1, 1, '"market","date","listed"'],
    ['di_soph_ssn_median_sm_share', 'softmax', 'lambda', 2.0, 30.0, 11.021992, FM, 1, 1],
    ['di_soph_ssn_median_sm_share', 'softmax', 'avg', 0.0, 1.5, False, FM, 1, 1],
    ['di_soph_ssn_median_sm_share', 'softmax', 'std', None, None, 0.047916, FM, 1, 1, None, None, None, 1,
     '''ds.loc[(ds["di_soph_ssn_median_sm_share"] > 0.0) & (ds["listed"].isin([2])), "di_soph_ssn_median_sm_share"].std()'''],

    # === Price: price_ret ===
    ['price_ret', 'mar', 'n', 10.0, 18.0, 17, 'int', 0, 1, '"bs_key"'],
    ['price_ret', 'mar', 'if_scale', None, None, False, 'bool', 1, 1, '"bs_key"'],
    ['price_ret', 'price', 'share_threshold', 0.0, 0.075, 0.001, FM, 0, 1],
    ['price_ret', 'price', 'threshold_var', None, None, '"units_base_share"', 'str', 1, 1],
    ['price_ret', 'price', 'weighted_by', None, None, '"di_soph_ssn_median"', 'str', 1, 1],
    ['price_ret', 'price', 'index_vars_dim', None, None, '"market"', 'str', 1, 1],
    ['price_ret', 'price', 'index_vars_time', None, None, '"date"', 'str', 1, 1],
    ['price_ret', 'price', 'year_start', None, None, 2021, 'int', 1, 1],
    ['price_ret', 'price', 'if_clutter', None, None, False, 'bool', 1, 1],
    ['price_ret', 'price', 'inn_coef', 0.0, 3.0, False, FM, 1, 1],
    ['price_ret_rel_market', 'mar', 'n', 4.0, 36.0, 15.835636, 'int', 0, 1, '"bs_key"'],
    ['price_ret_rel_market', 'mar', 'if_scale', None, None, True, 'bool', 1, 1, '"bs_key"'],
    ['price_ret_rel_market_mar', 'softmax', 'lambda', 3.0, 30.0, 20.207105, FM, 0, 1],
    ['price_ret_rel_market_mar', 'softmax', 'avg', -2.0, 409.8833, 204.94166, FM, 0, 1],
    ['price_ret_rel_market_mar', 'softmax', 'std', None, None, 105.030212, FM, 1, 1, None, None, None, 1,
     '''ds.loc[(ds["price_ret_rel_market_mar"] > 0.0) & (ds["listed"].isin([2])), "price_ret_rel_market_mar"].std()'''],

    # === adhoc_unparam_fin ===
    ['adhoc_unparam_fin', 'code_py.adhoc.model.adhoc_unparam_fin', 'dummy', None, None, False, 'bool', 1, 1],

    # === price_ret_volat_rel ===
    ['price_ret_volat_rel', 'softmax', 'lambda', 2.0, 30.0, 28.918645, FM, 0, 1],
    ['price_ret_volat_rel', 'softmax', 'avg', -1.5, 1.5, 0.749005, FM, 0, 1],
    ['price_ret_volat_rel', 'softmax', 'std', None, None, 0.530907, FM, 1, 1, None, None, None, 1,
     '''ds.loc[(ds["price_ret_volat_rel"] > 0.0) & (ds["listed"].isin([2])), "price_ret_volat_rel"].std()'''],

    # === ABA lag (NEW for p03) ===
    ['aba_lag', 'softmax', 'lambda', 3.0, 30.0, 15.0, FM, 0, 1],
    ['aba_lag', 'softmax', 'avg', 0.0, 100.0, aba_lag_avg, FM, 0, 1],
    ['aba_lag', 'softmax', 'std', None, None, aba_lag_std, FM, 1, 1],

    # === TV softmax (BHT input) ===
    ['brnd_tv_ots_clip', 'softmax', 'lambda', 5.0, 20.0, 9.0, FM, 0, 1],
    ['brnd_tv_ots_clip', 'softmax', 'avg', 0.0, avg_upper(tv_avg), tv_avg, FM, 0, 1],
    ['brnd_tv_ots_clip', 'softmax', 'std', None, None, tv_std, FM, 1, 1],

    # === OOH softmax (BHT input) ===
    ['brnd_ots_ooh', 'softmax', 'lambda', 3.0, 30.0, 5.0, FM, 0, 1],
    ['brnd_ots_ooh', 'softmax', 'avg', 0.0, avg_upper(ooh_avg), ooh_avg, FM, 0, 1],
    ['brnd_ots_ooh', 'softmax', 'std', None, None, ooh_std, FM, 1, 1],

    # === Coefs (LHS) ===
    [None, 'coef', 'tt_pk_tr_ssn', 0.1, 7.0, 1.5, FM, 0, 1],
    [None, 'coef', 'hdays', -0.1, 0.0, -0.0394, FM, 0, 1],
    [None, 'coef', 'price_volat_rel', -0.99, 0.0, -0.111989, FM, 0, 1],
    [None, 'coef', 'price', -0.99, 0.0, -0.2, FM, 0, 1],

    # === adhoc_bht_mult ===
    ['adhoc_mult', 'code_py.adhoc.model.adhoc_bht_mult', 'offset', 0.0, 1.0, 0.3, FM, 0, 1],

    # === TV: softmax_bht → lag → adstockp ===
    ['brnd_tv_ots_clip_softmax_bht', 'lag', 'n', 0.0, 3.0, 1, 'int', 0, 1, '"superbrand"', 'max'],
    ['brnd_tv_ots_clip_softmax_bht_lag', 'adstockp', 'a', 0.15, 0.95, 0.321491, FM, 0, 1, '"superbrand"', 'max'],

    # === OOH: softmax_bht → lag → adstockp ===
    ['brnd_ots_ooh_softmax_bht', 'lag', 'n', 0.0, 4.0, 0.368283, 'int', 0, 1, '"superbrand"', 'max'],
    ['brnd_ots_ooh_softmax_bht_lag', 'adstockp', 'a', 0.0, 0.95, 0.484365, FM, 0, 1, '"superbrand"', 'max'],

    # === Digital Performance ===
    ['brnd_cl_dgtind_perform', 'lag', 'n', 0.0, 3.0, 0, 'int', 0, 1, '"superbrand"', 'max'],
    ['brnd_cl_dgtind_perform_lag', 'softmax', 'lambda', 6.0, 30.0, 15.0, FM, 0, 1],
    ['brnd_cl_dgtind_perform_lag', 'softmax', 'avg', 0.0, avg_upper(dgt_perf_avg), dgt_perf_avg, FM, 0, 1],
    ['brnd_cl_dgtind_perform_lag', 'softmax', 'std', None, None, dgt_perf_std, FM, 1, 1],
    ['brnd_cl_dgtind_perform_lag_softmax', 'adstockp', 'a', 0.0, 0.25, 0.135582, FM, 0, 1, '"superbrand"', 'max'],

    # === Digital Display ===
    ['brnd_cl_dgtind_display', 'lag', 'n', 0.0, 2.0, 0, 'int', 0, 1, '"superbrand"', 'max'],
    ['brnd_cl_dgtind_display_lag', 'softmax', 'lambda', 4.0, 30.0, 5.0, FM, 0, 1],
    ['brnd_cl_dgtind_display_lag', 'softmax', 'avg', 0.0, avg_upper(dgt_disp_avg), dgt_disp_avg, FM, 0, 1],
    ['brnd_cl_dgtind_display_lag', 'softmax', 'std', None, None, dgt_disp_std, FM, 1, 1],
    ['brnd_cl_dgtind_display_lag_softmax', 'adstockp', 'a', 0.0, 0.85, 0.18, FM, 0, 1, '"superbrand"', 'max'],

    # === Digital Video ===
    ['brnd_cl_dgtind_video', 'lag', 'n', 0.0, 2.0, 0, 'int', 0, 1, '"superbrand"', 'max'],
    ['brnd_cl_dgtind_video_lag', 'softmax', 'lambda', 3.0, 30.0, 5.0, FM, 0, 1],
    ['brnd_cl_dgtind_video_lag', 'softmax', 'avg', 0.0, avg_upper(dgt_vid_avg), dgt_vid_avg, FM, 0, 1],
    ['brnd_cl_dgtind_video_lag', 'softmax', 'std', None, None, dgt_vid_std, FM, 1, 1],
    ['brnd_cl_dgtind_video_lag_softmax', 'adstockp', 'a', 0.0, 0.95, 0.068242, FM, 0, 1, '"superbrand"', 'max'],

    # === Digital Special ===
    ['brnd_cl_dgtind_special', 'lag', 'n', 0.0, 3.0, 0.100443, 'int', 0, 1, '"superbrand"', 'max'],
    ['brnd_cl_dgtind_special_lag', 'softmax', 'lambda', 3.0, 30.0, 5.0, FM, 0, 1],
    ['brnd_cl_dgtind_special_lag', 'softmax', 'avg', 0.0, avg_upper(dgt_spec_avg), dgt_spec_avg, FM, 0, 1],
    ['brnd_cl_dgtind_special_lag', 'softmax', 'std', None, None, dgt_spec_std, FM, 1, 1],
    ['brnd_cl_dgtind_special_lag_softmax', 'adstockp', 'a', 0.0, 0.95, 0.000232, FM, 0, 1, '"superbrand"', 'max'],

    # === Digital Indoor ===
    ['brnd_cl_dgtind_indoor', 'lag', 'n', 0.0, 3.0, 3, 'int', 0, 1, '"superbrand"', 'max'],
    ['brnd_cl_dgtind_indoor_lag', 'softmax', 'lambda', 2.0, 30.0, 5.0, FM, 0, 1],
    ['brnd_cl_dgtind_indoor_lag', 'softmax', 'avg', 0.0, avg_upper(dgt_indoor_avg), dgt_indoor_avg, FM, 0, 1],
    ['brnd_cl_dgtind_indoor_lag', 'softmax', 'std', None, None, dgt_indoor_std, FM, 1, 1],
    ['brnd_cl_dgtind_indoor_lag_softmax', 'adstockp', 'a', 0.0, 0.95, 0.8, FM, 0, 1, '"superbrand"', 'max'],

    # === E-com ===
    ['e_com', 'infl', 'infl', 0.9, 0.95, 0.95, FM, 1, 1],
    ['e_com', 'infl', 'year_start', None, None, 2021, FM, 1, 1],
    ['e_com_infl', 'lag', 'n', 0.0, 3.0, 2.967409, 'int', 1, 1, '"superbrand"', 'max'],
    ['e_com_infl_lag', 'softmax', 'lambda', 2.0, 30.0, 5.0, FM, 1, 1],
    ['e_com_infl_lag', 'softmax', 'avg', 0.0, avg_upper(ecom_avg), ecom_avg, FM, 1, 1],
    ['e_com_infl_lag', 'softmax', 'std', None, None, ecom_std, FM, 1, 1, None, None, None, 1,
     '''ds.loc[(ds["e_com_infl"] > 0.0) & (ds["listed"].isin([2])), "e_com_infl"].std()'''],
    ['e_com_infl_lag_softmax', 'adstockp', 'a', 0.0, 0.85, 0.789203, FM, 1, 1, '"superbrand"', 'max'],

    # === POS ===
    ['pos', 'infl', 'infl', 0.85, 1.0, 0.9, FM, 1, 1],
    ['pos', 'infl', 'year_start', None, None, 2021, FM, 1, 1],
    ['pos_infl', 'lag', 'n', 0.0, 6.0, 2, 'int', 1, 1, '"superbrand"', 'max'],
    ['pos_infl_lag', 'softmax', 'lambda', 2.0, 30.0, 5.0, FM, 1, 1],
    ['pos_infl_lag', 'softmax', 'avg', 0.0, avg_upper(pos_avg), pos_avg, FM, 1, 1],
    ['pos_infl_lag', 'softmax', 'std', None, None, pos_std, FM, 1, 1, None, None, None, 1,
     '''ds.loc[(ds["pos_infl"] > 0.0) & (ds["listed"].isin([2])), "pos_infl"].std()'''],
    ['pos_infl_lag_softmax', 'adstockp', 'a', 0.0, 0.85, 0.1, FM, 1, 1, '"superbrand"', 'max'],

    # === PR costs SM ===
    ['pr_costs_sm', 'infl', 'infl', 0.85, 1.0, 0.85, FM, 1, 1],
    ['pr_costs_sm', 'infl', 'year_start', None, None, 2021, FM, 1, 1],
    ['pr_costs_sm_infl', 'lag', 'n', 0.0, 8.0, 2, 'int', 1, 1, '"superbrand"', 'max'],
    ['pr_costs_sm_infl_lag', 'softmax', 'lambda', 3.0, 30.0, 5.0, FM, 1, 1],
    ['pr_costs_sm_infl_lag', 'softmax', 'avg', 0.0, avg_upper(pr_avg), pr_avg, FM, 1, 1],
    ['pr_costs_sm_infl_lag', 'softmax', 'std', None, None, pr_std, FM, 1, 1, None, None, None, 1,
     '''ds.loc[(ds["pr_costs_sm_infl"] > 0.0) & (ds["listed"].isin([2])), "pr_costs_sm_infl"].std()'''],
    ['pr_costs_sm_infl_lag_softmax', 'adstockp', 'a', 0.0, 0.95, 0.319055, FM, 1, 1, '"superbrand"', 'max'],
]

# Pad each row to 14 columns
for r in params_rows:
    while len(r) < 14:
        r.append(None)

ws_params = wb.active
ws_params.title = 'params'
write_sheet(ws_params, params_headers, params_rows, [30, 55, 18, 10, 18, 18, 10, 8, 10, 25, 10, 15, 8, 100])

# ═══════════════════════════════════════════
# Sheet: RHS
# ═══════════════════════════════════════════
rhs_rows = [
    ['di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '01_Distribution'],
    ['di_soph_ssn_median_sm_share_softmax', 1, 1, 1, 1, 1, 2, '01_Distribution'],
    ['brnd_tv_ots_clip_softmax_bht_lag_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '06_TV'],
    ['brnd_ots_ooh_softmax_bht_lag_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '07_OOH'],
    ['brnd_cl_dgtind_perform_lag_softmax_adstockp', 1, 1, 1, 1, 1, 2, '10_Digital_Performance'],
    ['brnd_cl_dgtind_display_lag_softmax_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '11_Digital_DA'],
    ['brnd_cl_dgtind_video_lag_softmax_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '11_Digital_Video'],
    ['brnd_cl_dgtind_special_lag_softmax_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '11_Digital_Special projects'],
    ['brnd_cl_dgtind_indoor_lag_softmax_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '08_Digital_Indoor'],
    ['e_com_infl_lag_softmax_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '13_E-com'],
    ['pr_costs_sm_infl_lag_softmax_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '14_PR'],
    ['pos_infl_lag_softmax_adstockp:di_soph_ssn_median_sm_softmax', 1, 1, 1, 1, 1, 2, '14_POS'],
]

ws_rhs = wb.create_sheet('RHS')
write_sheet(ws_rhs, ['token', 'fixed', 'if_active', 'sign', 'sign_weight', 'signif_weight', 'signif_threshold', 'display_var'],
            rhs_rows, [80, 8, 10, 8, 14, 14, 18, 30])

# ═══════════════════════════════════════════
# Sheet: setup
# ═══════════════════════════════════════════
ws_setup = wb.create_sheet('setup')
write_sheet(ws_setup, ['key', 'value'], [
    ['Y_var', 'aba'],
    ['model_type', 'LME'],
    ['fixed_effect_var', 'bs_key'],
    ['target_superbrand', 'brand_1'],
    ['period_var', 'week'],
    ['price_var', 'price_ret'],
    ['conversion_var', None],
    ['subfolder_orig', '20260617'],
    ['custom_scoring', 'rmse_bs_key * ((sign_t_value_st + 1) ** 10)'],
    ['target_audience', 'all_20_45'],
    ['bs_key', 'superbrand'],
    ['bs_key', 'master'],
    ['bs_key', 'market'],
    ['lme_method', 'powell'],
    ['lme_method', 'lbfgs'],
], [20, 80])

# ═══════════════════════════════════════════
# Sheet: scoring
# ═══════════════════════════════════════════
ws_scoring = wb.create_sheet('scoring')
write_sheet(ws_scoring, ['metrics', 'if_invert', 'width', 'weight', 'if_active'], [
    ['rmse_bs_key', 0, 0.3, 0.166667, 1],
    ['r_squared', 1, 0.3, 1.0, 0],
    ['r_squared_adjusted', 0, 0.3, 1.0, 0],
    ['t_value', 1, 0.3, 0.5, 1],
    ['p_value', 0, 0.3, 0.166667, 1],
    ['sign_simple', 0, 0.3, 1.0, 0],
    ['sign_t_value_st', 0, 0.3, 0.166667, 1],
    ['vif', 0, 0.3, 1.0, 0],
    ['durbin_watson', 0, 0.3, 1.0, 0],
], [20, 10, 8, 10, 10])

# ═══════════════════════════════════════════
# Sheet: LHS
# ═══════════════════════════════════════════
ws_lhs = wb.create_sheet('LHS')
write_sheet(ws_lhs, ['type', 'name', 'token', 'if_active', 'display_var'], [
    ['multiplicative', 'ssn_mrk', '(coef.tt_pk_tr_ssn + 1) / (coef.tt_pk_tr_ssn * ds.tt_pk_tr_ssn  + 1)', 1, '20_seasonality'],
    ['multiplicative', 'hdays', '1 + ds.hdays * coef.hdays', 1, '20_seasonality'],
    ['multiplicative', 'price', '1 + coef.price * ds.price_ret_rel_market_mar_softmax', 1, '05_Relative price'],
    ['multiplicative', 'price_volat_rel', '1 + coef.price_volat_rel * ds.price_ret_volat_rel_softmax', 1, '06_Price volatility'],
], [18, 16, 80, 10, 25])

# ═══════════════════════════════════════════
# Sheet: trans_path_df
# ═══════════════════════════════════════════
ws_tpdf = wb.create_sheet('trans_path_df')
write_sheet(ws_tpdf, ['variable', 'variable_fin', 'price', 'display_var'], [
    ['brnd_tv_ots_clip', None, 100.0, None],
    ['brnd_ots_ooh', None, 100.0, None],
    ['brnd_cl_dgtind_perform', None, 10.0, None],
    ['brnd_cl_dgtind_display', None, 0.1, None],
    ['brnd_cl_dgtind_video', None, 0.3, None],
    ['brnd_cl_dgtind_special', None, 0.2, None],
    ['brnd_cl_dgtind_indoor', None, 0.3, None],
    ['e_com', None, 1.0, None],
    ['pr_costs_sm', None, 1.0, None],
    ['pos', None, 1.0, None],
    ['adhoc_mult', 'adhoc_mult_code_py.adhoc.model.adhoc_bht_mult', 1.0, 'adhoc_mult_code_py.adhoc.model.adhoc_bht_mult'],
], [50, 70, 10, 70])

wb.save(out_path)
print(f'Created: {out_path}')
